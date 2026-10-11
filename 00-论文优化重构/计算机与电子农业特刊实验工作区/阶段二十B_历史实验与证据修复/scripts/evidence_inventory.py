from __future__ import annotations

import csv
import json
import re
from pathlib import Path

from openpyxl import load_workbook

from common import DATA, ROOT, WORK, dimensions, git_origin, markdown_table, read_csv, relative, report, sha256, write_csv

TABLES = WORK / '04-结果表格模板'
PHENOTYPE = ROOT / '植株数据.xlsx'
RELEASE = ROOT / 'ForeSplat_real_release'
TEX = WORK / 'ForeSplat__Foreground_Object_2D_Gaussian_Splatting_via_Semantic_Mask_Guided_Task_Reformulation_for_Plant/ForeSplat.tex'


def file_evidence(path, evidence_type, claim, status, note):
    return {'evidence_type': evidence_type, 'path': relative(path), 'source_commit': git_origin(path) if path.exists() else 'MISSING', 'sha256': sha256(path) if path.exists() else 'UNKNOWN', 'claim': claim, 'status': status, 'note': note}


def scale_audit():
    evidence = []
    values = []
    workbook = load_workbook(PHENOTYPE, read_only=True, data_only=True)
    sheet = workbook.active
    headers = [str(value or '') for value in next(sheet.iter_rows(values_only=True))]
    rows = list(sheet.iter_rows(min_row=2, values_only=True))
    terms = ['pot', 'diameter', 'radius', 'scale', 'marker', '盆', '直径', '半径', '尺度', '标尺']
    matched_headers = [header for header in headers if any(term.lower() in header.lower() for term in terms)]
    evidence.append(file_evidence(PHENOTYPE, 'measurement_workbook', '21 plant rows and four manual/virtual trait families are present', 'VERIFIED', f'headers={headers}; scale_related_headers={matched_headers or "NONE"}'))
    evidence.append({'evidence_type': 'workbook_schema', 'path': relative(PHENOTYPE), 'source_commit': git_origin(PHENOTYPE), 'sha256': sha256(PHENOTYPE), 'claim': 'No pot diameter, scale marker or scale factor column', 'status': 'NOT FOUND', 'note': f'{len(rows)} data rows; {len(headers)} columns'})
    for candidate in [RELEASE / 'scripts/measure_traits_entry.py', RELEASE / 'configs/default_foresplat.yaml', RELEASE / 'configs/table8_phenotype_measurement.yaml', TEX]:
        text = candidate.read_text(encoding='utf-8', errors='replace') if candidate.exists() else ''
        hits = [line.strip() for line in text.splitlines() if re.search(r'pot.?diameter|pot.?radius|scale.?marker|scale.?factor|world.?scale|盆径|标尺|尺度', line, re.I)]
        evidence.append(file_evidence(candidate, 'implementation_or_manuscript', 'scale recovery or trait measurement path', 'MENTIONED' if hits else 'NO IMPLEMENTATION EVIDENCE', '\n'.join(hits[:8]) if hits else 'No matching line'))
    trait_csv = TABLES / '实验五_人工虚拟表型汇总.csv'
    pair_csv = TABLES / '实验五_掩膜到表型误差结果表.csv'
    evidence.append(file_evidence(trait_csv, 'summary_table', 'reported aggregate trait agreement', 'VERIFIED', 'Contains already aggregated values; not a 3D extraction implementation'))
    evidence.append(file_evidence(pair_csv, 'summary_table', 'reported paired trait error and mask-quality notes', 'VERIFIED', 'The row says values came from 植株数据.xlsx and does not identify a Mesh per trait'))
    for path in sorted((ROOT / '06-MESH-new').glob('*_mesh*.log')):
        text = path.read_text(encoding='utf-8', errors='replace')
        hits = [line.strip() for line in text.splitlines() if re.search(r'pot|scale|diameter|trait|height|width|测量|尺度|盆', line, re.I)]
        evidence.append(file_evidence(path, 'mesh_log', 'mesh log contains physical-scale or trait extraction trace', 'FOUND' if hits else 'NO TRACE', '\n'.join(hits[:5]) if hits else 'No pot/scale/trait fields'))
    report('SCALE_RECOVERY_AUDIT.md', f'''# 尺度恢复审计

`植株数据.xlsx` 是真实工作簿：1 个 sheet、{len(rows)} 个植株行、{len(headers)} 列；字段为人工株高/冠幅/三片叶的人工与虚拟数值。
工作簿没有盆径、半径、scale marker 或 scale factor 字段，不能由它恢复物理尺度。
退稿稿件在 Table 8 叙述中依赖物理尺度，但仓库中没有逐样本盆径、标尺图、标定系数或传递记录。
`measure_traits_entry.py` 只要求 `--scale-marker`，输出外部工具的输入清单，并明确完整交互/半自动测量工具未随轻量发布包提供；这不是尺度恢复实现。
`06-MESH-new/*_mesh.log` 共 {len(list((ROOT / '06-MESH-new').glob('*_mesh*.log')))} 个日志，没有 pot/scale/trait 记录（详见 CSV）。
因此当前状态是 `USER_INPUT_REQUIRED`，不测量、不猜盆径、不把工作簿统计脚本当成三维测量算法。
''')
    return evidence, rows, headers


def phenotype_audit(evidence, rows, headers):
    pair = read_csv(TABLES / '实验五_掩膜到表型误差结果表.csv')
    trait = read_csv(TABLES / '实验五_人工虚拟表型汇总.csv')
    source = [
        {'source': relative(PHENOTYPE), 'role': 'manual_and_virtual_trait_values', 'n_rows': len(rows), 'status': 'VERIFIED', 'mesh_id_per_value': 'NOT RECORDED', 'scale_factor_per_sample': 'NOT FOUND', 'measurement_method': 'UNKNOWN_EXTERNAL_OR_MANUAL_INTERACTIVE', 'evidence_type': 'measurement_workbook', 'source_commit': git_origin(PHENOTYPE), 'sha256': sha256(PHENOTYPE), 'claim': 'Manual and virtual trait values are present', 'note': 'Four trait families; no upstream mesh/scale linkage recorded'},
        {'source': relative(TABLES / '实验五_人工虚拟表型汇总.csv'), 'role': 'aggregate_error_statistics', 'n_rows': len(trait), 'status': 'VERIFIED_AS_STATISTICS_ONLY', 'mesh_id_per_value': 'NOT RECORDED', 'scale_factor_per_sample': 'NOT FOUND', 'measurement_method': 'Statistics over workbook values', 'evidence_type': 'summary_table', 'source_commit': git_origin(TABLES / '实验五_人工虚拟表型汇总.csv'), 'sha256': sha256(TABLES / '实验五_人工虚拟表型汇总.csv'), 'claim': 'Reported aggregate trait agreement', 'note': 'Already aggregated; not a 3D extraction implementation'},
        {'source': relative(TABLES / '实验五_掩膜到表型误差结果表.csv'), 'role': 'mask_quality_sensitivity_summary', 'n_rows': len(pair), 'status': 'PARTIAL', 'mesh_id_per_value': 'NOT RECORDED', 'scale_factor_per_sample': 'NOT FOUND', 'measurement_method': 'Summary row; no executable 3D extraction', 'evidence_type': 'summary_table', 'source_commit': git_origin(TABLES / '实验五_掩膜到表型误差结果表.csv'), 'sha256': sha256(TABLES / '实验五_掩膜到表型误差结果表.csv'), 'claim': 'Reported mask-quality sensitivity summary', 'note': 'Does not identify a Mesh per trait'},
    ]
    for item in evidence:
        source.append({
            'source': item['path'],
            'role': item['evidence_type'],
            'n_rows': '',
            'status': item['status'],
            'mesh_id_per_value': 'NOT RECORDED',
            'scale_factor_per_sample': 'NOT FOUND',
            'measurement_method': 'NOT ESTABLISHED',
            'evidence_type': item['evidence_type'],
            'source_commit': item['source_commit'],
            'sha256': item['sha256'],
            'claim': item['claim'],
            'note': item['note'],
        })
    write_csv('PHENOTYPE_DATA_SOURCE_MATRIX.csv', source)
    report('PHENOTYPE_PROVENANCE_AUDIT.md', '''# 表型数据来源审计

人工与虚拟值可追溯到 `植株数据.xlsx`，并可重算 Table 8 的均值、MAE、RMSE、MAPE 和 R²；这些值是已记录结果，不是本仓库现场从 Mesh 提取的结果。
`生成实验一五当前结果.py` 读取工作簿并生成统计表，没有点云/Mesh读取、尺度变换、点选、叶片 landmark 或 trait geometry 函数。
`measure_traits_entry.py` 只写测量 manifest，要求外部 `mesh`、`scale-marker` 和可选 `landmarks`；发布版说明完整工具是交互/半自动且未提供。
因此四项性状的统计结果可以作为 `LEGACY_RECORDED_STATISTICS`，但不能标为 `REPRODUCIBLE_3D_EXTRACTION`；每个虚拟值对应的 Mesh、坐标系、尺度系数和点选记录未找到。
Table 8 的 n=21（株高/冠幅）和 n=63（叶长/叶宽）来自工作簿结构，但不证明所有样本都有可复核三维产物。
''')
    report('USER_MEASUREMENT_INPUT_REQUIRED.md', '''# 用户测量资料清单

以下资料是完成尺度和表型完全复现真正需要用户补充的最小集合；本阶段不补造数字：

- 21 个植株样本逐样本真实盆径（单位、测量工具、测量时点和重复测量/不确定度）。
- 每个虚拟性状值对应的 Mesh 文件版本、样本 ID、Git/运行 ID 和坐标系。
- 尺度恢复路径：盆径在 RGB、COLMAP、点云还是 Mesh 中标定，以及像素/世界单位转换系数。
- 每个样本的虚拟测量操作记录：株高/冠幅端点、代表叶 ID、叶长/叶宽 landmark 或交互标注文件。
- 21 个样本中缺失的真实 Mesh、尺度 marker 图或 scale-factor 日志。
- 若表型数字并非该仓库 Mesh 生成：外部工具名称、版本、命令和导出 CSV 的原始文件。

已确认不需要用户补充：`植株数据.xlsx` 本身已在仓库中，且四项性状的人工/虚拟统计汇总可读；缺的是上游物理尺度和三维操作证据。
''')


if __name__ == '__main__':
    evidence, rows, headers = scale_audit()
    phenotype_audit(evidence, rows, headers)
    print('scale/phenotype inventory complete')
